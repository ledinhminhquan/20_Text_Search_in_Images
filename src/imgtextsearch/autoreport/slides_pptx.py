"""Generate the submission slides.pptx (python-pptx) - ~12 concise slides for imgtextsearch.
Degrades to a Markdown outline if python-pptx is unavailable.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from ..config import AppConfig, run_dir
from ..logging_utils import get_logger
from . import charts as charts_mod
from .artifact_loader import exact_match, load_artifacts, ocr_cer, sys_recall

logger = get_logger(__name__)


def _pct(v):
    return f"{v*100:.1f}%" if isinstance(v, (int, float)) and not isinstance(v, bool) else "?"


def _slides(cfg: AppConfig, arts: Dict[str, Any]) -> List[Tuple[str, List[str]]]:
    r1 = sys_recall(arts, "hybrid", "@1")
    rr = sys_recall(arts, "random", "@1")
    ex = exact_match(arts)
    res = (f"hybrid Recall@1 {_pct(r1)} vs random {_pct(rr)}" if r1 is not None
           else "train + evaluate to populate results")
    exline = (f"exact-match precision {_pct(ex)}; OCR CER {ocr_cer(arts):.3f}" if ex is not None and ocr_cer(arts) is not None
              else "exact-match precision + OCR CER")
    return [
        ("Text Search within Images",
         [f"{cfg.author} - Student {cfg.student_id}", "NLP in Industry - Final Assignment",
          "Search a collection by the TEXT inside the images (OCR)",
          "OCR -> text index (BM25 + dense -> RRF) -> verify -> snippet",
          "Abstains ('not found') when no image contains the text"]),
        ("Business Problem & Motivation",
         ["Document/receipt/invoice archives, scanned-contract search, e-discovery",
          "Find images by the WORDS they contain - not by visual content (that was P19)",
          "The fix = OCR + a text index + literal verification + an agent",
          "Only the dense text retriever is trained; OCR + BM25 are algorithmic"]),
        ("Proposed Solution",
         ["OCR each image -> the per-image text is the searchable 'document'",
          "Hybrid index: BM25 (exact terms) + dense (semantic) -> RRF",
          "Verify the query term LITERALLY appears (fuzzy to OCR errors) + snippet",
          "Abstain when no image contains the query text"]),
        ("System Architecture",
         ["images -> OCR (Tesseract / SeedEngine) -> text index",
          "query: parse (D1) -> search (D2) -> coverage (D3)",
          "-> verify + snippet (D4) -> abstain (D5)",
          "Runs fully offline (BM25 + SeedEngine, numpy fallback) for tests/CI"]),
        ("Data (Scene-text + Synthetic)",
         ["MiXaiLL76/TextOCR_OCR (MIT, image+text), IIIT5K_OCR; cord-v2 receipts (CC-BY)",
          "OCR text source PleIAs (CC0); docvqa_1200 (query+gold-span, license-flag)",
          "rvl_cdip/funsd images (NC/research-flag) for scale tests",
          "NO 'find image containing X' benchmark -> synthetic rendered-text generator"]),
        ("The OCR Front-End + Retriever",
         ["Tesseract (Apache) default; trocr-base-printed (MIT) neural upgrade",
          "Dense retriever bge-small-en-v1.5 (MIT), fine-tune MNRL on (query, OCR-text)",
          "BM25 = the exact-term arm (critical for literal search); RRF fusion",
          "Baselines: BM25-only, dense-only, random"]),
        ("Metrics",
         ["Text-in-image Recall@1/5/10 + MRR + median rank (multi-gold)",
          "OCR CER/WER (OCR vs gold text)",
          "Exact-match precision (returned images literally contain the term)",
          "The hybrid combines exact + semantic; fuzzy verify recovers OCR noise"]),
        ("The 5-Decision Agent",
         ["D1 query gate+classify - D2 BM25+dense->RRF search",
          "D3 coverage/confidence gate",
          "D4 snippet + LITERAL-term verify (fuzzy to OCR errors)",
          "D5 ABSTAIN 'not found' when no image contains the text"]),
        ("Evaluation Results",
         [res, exline,
          "Recall@k + MRR + OCR CER + exact-match, vs BM25-only / dense-only / random",
          "Literal verification + snippet + abstention = the OCR-search value-add"]),
        ("Deployment Overview",
         ["FastAPI POST /search (query -> images + scores + snippet + exact-match flag) + /healthz",
          "Gradio demo (type a query -> see matching images + highlighted snippets)",
          "Docker (tesseract-ocr + fonts + libGL) + HF Space; BM25/SeedEngine offline fallback",
          "OCR-index built once at startup; metadata-only job logging"]),
        ("Continual Learning, Monitoring & Ethics",
         ["Index freshness: new images -> incremental OCR + re-index",
          "monitor-log: abstain/exact-match rate + top-score drift + latency",
          "Privacy: scanned docs = highly sensitive PII -> access control, redaction, no retention",
          "OCR-quality bias -> unequal recall by script/quality; abstain + snippet for review"]),
        ("Key Takeaways & Future Work",
         ["A literal-verifying, snippet-showing, abstaining OCR-search pipeline",
          "BM25 exact + dense semantic + fuzzy verify beat blind semantic ranking",
          "Future: neural OCR (TrOCR), phrase/proximity search, multilingual scripts",
          "Future: PII redaction, ANN at scale, post-OCR correction before indexing"]),
    ]


def generate_slides(cfg: AppConfig, title: Optional[str] = None, author: Optional[str] = None,
                    out_path: Optional[str] = None) -> str:
    arts = load_artifacts(cfg)
    out_path = Path(out_path) if out_path else run_dir() / "report" / "slides.pptx"
    out_path.parent.mkdir(parents=True, exist_ok=True)
    slides = _slides(cfg, arts)
    try:
        from pptx import Presentation
        from pptx.dml.color import RGBColor
        from pptx.util import Inches, Pt
    except Exception as exc:
        logger.warning("python-pptx unavailable (%s); writing markdown outline", exc)
        md = "\n\n".join(f"## {t}\n" + "\n".join(f"- {b}" for b in bs) for t, bs in slides)
        alt = out_path.with_suffix(".md")
        alt.write_text(md, encoding="utf-8")
        return str(alt)

    try:
        chart = charts_mod.recall_chart(arts, run_dir() / "report" / "slide_recall.png")
    except Exception:
        chart = None
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    accent = RGBColor(0x2B, 0x6C, 0xB0)
    for i, (t, bullets) in enumerate(slides):
        slide = prs.slides.add_slide(blank)
        bar = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(13.333), Inches(1.1))
        bar.fill.solid(); bar.fill.fore_color.rgb = accent; bar.line.fill.background()
        tf = bar.text_frame; tf.text = t
        tf.paragraphs[0].font.size = Pt(28); tf.paragraphs[0].font.bold = True
        tf.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255)
        body = slide.shapes.add_textbox(Inches(0.6), Inches(1.5),
                                        Inches(8.3 if (i == 8 and chart) else 12), Inches(5.4))
        bt = body.text_frame; bt.word_wrap = True
        for j, bp in enumerate(bullets):
            p = bt.paragraphs[0] if j == 0 else bt.add_paragraph()
            p.text = "-  " + bp; p.font.size = Pt(20); p.space_after = Pt(10)
        if i == 8 and chart:
            slide.shapes.add_picture(str(chart), Inches(8.9), Inches(1.7), width=Inches(4.0))
        foot = slide.shapes.add_textbox(Inches(0.4), Inches(7.0), Inches(12.5), Inches(0.4))
        foot.text_frame.text = f"{title or cfg.project_title} - {author or cfg.author} ({cfg.student_id})"
        foot.text_frame.paragraphs[0].font.size = Pt(9)
    prs.save(str(out_path))
    logger.info("Slides -> %s", out_path)
    return str(out_path)


__all__ = ["generate_slides"]
